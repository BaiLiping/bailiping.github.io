#!/usr/bin/env python3
"""Export static PDF and screenshot-based PowerPoint with presenter notes.
Requires Node Playwright and python-pptx. Set CHROMIUM_EXECUTABLE for a system browser.
The HTML is the interactive master; static lab explanations remain in every export.
"""
import argparse, json, os, shutil, tempfile, subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1]

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'exports')
    parser.add_argument('--screenshots',type=Path,help='Optional persistent slide PNG directory')
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    deck=json.loads((ROOT/'src/deck.json').read_text())
    with tempfile.TemporaryDirectory() as temp:
        screenshots=args.screenshots or Path(temp);screenshots.mkdir(parents=True,exist_ok=True)
        subprocess.run(['node',str(ROOT/'tools/export.cjs'),'--output',str(args.output),'--screenshots',str(screenshots)],check=True)
        rendered=json.loads((args.output/'bento-slides.json').read_text())
        presentation=Presentation();presentation.slide_width=Inches(13.333333);presentation.slide_height=Inches(7.5)
        presentation.core_properties.title='Cramér–Rao Bound: Precision Has a Floor'
        presentation.core_properties.subject='23 slides with three companion interactive labs'
        presentation.core_properties.author='Bai Liping'
        presentation.core_properties.keywords='Cramer Rao; Fisher information; estimation; localization; bias'
        presentation.core_properties.comments='Static slide images. The companion HTML contains the live controls; source text and formulas are included in the site package.'
        sources={
          'S1':'Stanford STATS 200, Lecture 15: https://web.stanford.edu/class/archive/stats/stats200/stats200.1172/Lecture15.pdf',
          'S2':'Ly et al. (2017), A Tutorial on Fisher Information: https://arxiv.org/abs/1705.01064',
          'S3':'Shen & Win (2010), Fundamental Limits of Wideband Localization—Part I: https://arxiv.org/abs/1006.0888',
          'S4':'Stanford STATS 200, Lecture 14: https://web.stanford.edu/class/archive/stats/stats200/stats200.1172/Lecture14.pdf'}
        for i,item in enumerate(deck):
            slide=presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(screenshots/f'slide-{i+1:02d}.png'),0,0,width=presentation.slide_width,height=presentation.slide_height)
            description=slide.shapes[0]._element.xpath('.//p:cNvPr')[0]
            description.set('descr',f'Slide {i+1}: {item["title"]}. {item["subtitle"]}. See speaker notes and the HTML master for the full content.')
            note=f'{i+1:02d}. {item["title"]}\n\n{item["subtitle"]}\n\n{rendered[i]["notes"]}\n\n'
            if item['id'].endswith('-lab'):
                note+='LIVE LAB: Open Cramer-Rao-Bound.html in a browser and choose TRY LIVE on this slide. This PowerPoint is a static snapshot.\n\n'
            note+='Sources\n'+'\n'.join(sources[key] for key in item['sources'])
            slide.notes_slide.notes_text_frame.text=note
        presentation.save(args.output/'Cramer-Rao-Bound.pptx')
    shutil.copyfile(ROOT/'index.html',args.output/'Cramer-Rao-Bound.html')
    for name in ['live','guide']:
        shutil.copytree(ROOT/name,args.output/name,dirs_exist_ok=True)
    print(f'Exported HTML, PDF and PowerPoint to {args.output}')

if __name__=='__main__':main()
